import argparse
import csv
import json
import sys
import time
import shutil
from pathlib import Path
import re
from io import BytesIO
from zipfile import is_zipfile

AI_CHECK_SETTINGS_PATH = Path(__file__).with_name("ai_check_settings.json")


def load_ai_instructions(path=None):
    """Read shared checker instructions fresh for every run."""
    config = json.loads(Path(path or AI_CHECK_SETTINGS_PATH).read_text(encoding="utf-8-sig"))
    if not isinstance(config, dict) or config.get("schema_version") != 1:
        raise ValueError("AI check settings must have schema_version 1.")
    lines = config.get("instruction_lines")
    if not isinstance(lines, list) or not lines or any(not isinstance(line, str) or not line.strip() for line in lines):
        raise ValueError("instruction_lines must be a nonempty list of nonempty strings.")
    return lines


def cache_directory():
    """Use one persistent product cache beside the tool for all projects."""
    return Path(__file__).resolve().parent / "lcsc"


def import_local_cache(local_directory):
    """Copy missing legacy product files without replacing shared or local data."""
    source_directory = Path(local_directory)
    destination_directory = cache_directory()
    destination_directory.mkdir(parents=True, exist_ok=True)
    if not source_directory.is_dir() or source_directory.resolve() == destination_directory.resolve():
        return 0, 0
    copied = skipped = 0
    for source in sorted(source_directory.iterdir()):
        if not source.is_file() or not re.fullmatch(r"C\d+\.(?:htm|json)", source.name, re.I):
            continue
        destination = destination_directory / f"{source.stem.upper()}{source.suffix.lower()}"
        created = False
        try:
            with source.open("rb") as input_file:
                with destination.open("xb") as output_file:
                    created = True
                    shutil.copyfileobj(input_file, output_file)
        except FileExistsError:
            skipped += 1
            continue
        except Exception:
            if created:
                destination.unlink()
            raise
        copied += 1
    return copied, skipped


def _read_csv_parts(path):
    parts = []
    with open(path, newline="", encoding="utf-8") as f:
        sample = f.read(4096)
        f.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample)
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(f, dialect)
        rows = list(reader)
    if not rows:
        return parts
    for row in rows[1:]:
        if not row:
            continue
        parts.append(row[-1])
    return parts


def _read_csv_rows(path, with_designators=False):
    rows_out = []
    with open(path, newline="", encoding="utf-8") as f:
        sample = f.read(4096)
        f.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample)
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(f, dialect)
        rows = list(reader)
    if not rows:
        return rows_out
    ref_column = next((i for i, header in enumerate(rows[0])
                       if header.strip().casefold() in {"designator", "designators", "refdes"}), None)
    for row in rows[1:]:
        if not row:
            continue
        desc = row[0] if len(row) > 0 else ""
        part = row[-1] if len(row) > 0 else ""
        refs = row[ref_column] if ref_column is not None and ref_column < len(row) else ""
        rows_out.append((desc, part, refs) if with_designators else (desc, part))
    return rows_out


def _read_excel_parts(path):
    return [part for _description, part in _read_excel_rows(path)]


def _read_excel_rows(path, with_designators=False):
    try:
        import pandas as pd
    except Exception as exc:
        raise RuntimeError(
            "pandas is required for .xls/.xlsx parsing. Install dependencies from requirements.txt."
        ) from exc

    if is_zipfile(path):
        # Some JLCPCB exports are OOXML workbooks saved with a .xls extension.
        # Load by content, and reset dimensions because these exports can report A1
        # even when the sheet contains many rows and columns.
        try:
            from openpyxl import load_workbook
        except Exception as exc:
            raise RuntimeError(
                "openpyxl is required for .xlsx workbook parsing. Install dependencies from requirements.txt."
            ) from exc
        workbook = load_workbook(
            BytesIO(Path(path).read_bytes()), data_only=True, read_only=True
        )
        sheet = workbook.active
        sheet.reset_dimensions()
        grid = [list(row) for row in sheet.iter_rows(values_only=True)]
        workbook.close()
    else:
        ext = Path(path).suffix.lower()
        if ext == ".xls":
            df = pd.read_excel(path, engine="xlrd")
        else:
            df = pd.read_excel(path, engine="openpyxl")
        if df.empty:
            return []
        grid = [list(df.columns)] + list(df.itertuples(index=False, name=None))

    def cell_text(value):
        if value is None or pd.isna(value):
            return ""
        return str(value).strip()

    # JLCPCB BOM Matching exports have a title block above a table. Use its
    # Comment column and the last JLCPCB Part # column (the matched part; the
    # earlier same-named column is the user's optional input).
    matching_header = None
    for row_index, row in enumerate(grid[:50]):
        headers = [cell_text(value) for value in row]
        lowered = [header.casefold() for header in headers]
        if (
            "comment" in lowered
            and "designator" in lowered
            and "matches" in lowered
            and "product link" in lowered
        ):
            part_columns = [
                index for index, header in enumerate(lowered)
                if header.startswith("jlcpcb part #")
            ]
            if part_columns:
                matching_header = (row_index, lowered.index("comment"), part_columns[-1])
                break

    if matching_header:
        header_row, description_column, part_column = matching_header
        ref_column = [cell_text(value).casefold() for value in grid[header_row]].index("designator")
        rows_out = []
        for row in grid[header_row + 1 :]:
            desc = cell_text(row[description_column]) if description_column < len(row) else ""
            part = cell_text(row[part_column]) if part_column < len(row) else ""
            if desc or part:
                refs = cell_text(row[ref_column]) if ref_column < len(row) else ""
                rows_out.append((desc, part, refs) if with_designators else (desc, part))
        return rows_out

    # Original tool format: first row is headers; descriptions are in the first
    # column and LCSC C-part numbers are in the final column.
    if not grid:
        return []
    ref_column = next((i for i, header in enumerate(grid[0])
                       if cell_text(header).casefold() in {"designator", "designators", "refdes"}), None)
    rows_out = []
    for row in grid[1:]:
        if not row:
            continue
        desc = cell_text(row[0]) if len(row) > 0 else ""
        part = cell_text(row[-1]) if row else ""
        if desc or part:
            refs = cell_text(row[ref_column]) if ref_column is not None and ref_column < len(row) else ""
            rows_out.append((desc, part, refs) if with_designators else (desc, part))
    return rows_out


def load_check_notes(path):
    """Read saved project rows, or the editor's temporary in-memory snapshot."""
    from bom_editor_core import expand_rows
    document = json.loads(Path(path).read_text(encoding="utf-8-sig"))
    return expand_rows(document["rows"])


def project_check_rows(records):
    """Use the same selected parts, grouping and normalization as Generate."""
    from bom_editor_core import FIELDS, grouped_rows, selected_component
    from bom_conversions import assembly_bom_rows, TARGET_COLUMNS
    selected = {ref: {field: selected_component(record)[field] for field in FIELDS}
                for ref, record in records.items()}
    rows = []
    for row in assembly_bom_rows(grouped_rows(selected)):
        description = f"{row['Designator']}: {row['Comment']}"
        if row["Footprint"]:
            description += f" [footprint: {row['Footprint']}]"
        rows.append((description, row[TARGET_COLUMNS[3]], row["Designator"]))
    return rows


def notes_for_row(records, designators, part):
    from bom_editor_core import REFS, selected_component
    refs = REFS.findall(str(designators).upper())
    if refs:
        candidates = [(ref, records.get(ref, {})) for ref in refs]
    else:
        # Legacy workbooks may lack designators. Keep per-part labels in that case.
        candidates = [(ref, row) for ref, row in records.items()
                      if selected_component(row).get("LCSC Part #", "").upper() == part.upper()]
    notes = []
    for ref, row in candidates:
        texts = [row.get("Notes", "")]
        if row.get("Output selection", 0):
            texts.append(selected_component(row).get("Notes", ""))
        texts = list(dict.fromkeys(" ".join(text.split()) for text in texts if text.strip()))
        notes.append((ref, "; ".join(texts)))
    if notes and refs and len({note for _, note in notes}) == 1:
        return notes[0][1]
    return "; ".join(f"{ref}: {note}" for ref, note in notes if note)


def verification_line(description, model, name, notes=""):
    line = f"{description}\t{model}: {name}"
    return line + (f" // {notes}" if notes else "") + "\n"


def _normalize_parts(values):
    parts = []
    for v in values:
        if v is None:
            continue
        s = str(v).strip()
        if not s:
            continue
        # Only include LCSC C-part numbers
        if not s.upper().startswith("C"):
            continue
        parts.append(s.upper())
    return sorted(set(parts))


def _download_html(part, out_path, sleep_s=0.2, overwrite=False):
    import requests

    if out_path.exists() and not overwrite:
        return False

    print(f"{part}: downloading...")
    url = f"https://www.lcsc.com/product-detail/{part}.html"
    headers = {
        "User-Agent": "Mozilla/5.0 (compatible; JLCPCB_BOM_Downloader/1.0)",
        "Accept": "text/html,application/xhtml+xml",
    }
    resp = requests.get(url, headers=headers, timeout=30)
    resp.raise_for_status()
    content = resp.content
    text = None
    try:
        text = resp.text
    except Exception:
        text = None

    # If LCSC returns "Page Not Found", try JLCPCB part detail as a fallback.
    if text and "Page Not Found | LCSC Electronics" in text:
        print(f"{part}: LCSC not found, trying JLCPCB...")
        jlc_url = f"https://jlcpcb.com/partdetail/{part}"
        jlc_headers = {
            "User-Agent": "Mozilla/5.0 (compatible; JLCPCB_BOM_Downloader/1.0)",
            "Accept": "text/html,application/xhtml+xml",
        }
        jlc_resp = requests.get(jlc_url, headers=jlc_headers, timeout=30)
        jlc_resp.raise_for_status()
        content = jlc_resp.content

    out_path.write_bytes(content)
    time.sleep(sleep_s)
    return True


def _extract_nuxt_blob(html_text):
    marker = "window.__NUXT__="
    start = html_text.find(marker)
    if start == -1:
        return None
    end = html_text.find("</script>", start)
    if end == -1:
        return None
    snippet = html_text[start:end]
    # Trim any trailing semicolons/spaces
    return snippet.strip().rstrip(";")


def _extract_next_data(html_text):
    match = re.search(
        r'<script\b[^>]*id=["\']__NEXT_DATA__["\'][^>]*>(.*?)</script>',
        html_text,
        flags=re.IGNORECASE | re.DOTALL,
    )
    if not match:
        return None
    try:
        return json.loads(match.group(1))
    except json.JSONDecodeError:
        return None


def _parse_next_detail(html_text):
    next_data = _extract_next_data(html_text)
    if not isinstance(next_data, dict):
        return None
    page_props = next_data.get("props", {}).get("pageProps", {})
    detail = page_props.get("webData") if isinstance(page_props, dict) else None
    if isinstance(detail, dict) and detail.get("productCode"):
        return {"detail": detail, "breadcrumbs": None, "source": "lcsc"}
    return None


def _parse_nuxt_detail(html_text):
    try:
        import quickjs  # type: ignore
    except Exception as exc:
        raise RuntimeError(
            "quickjs is required for HTML scraping. Install dependencies from requirements.txt."
        ) from exc

    snippet = _extract_nuxt_blob(html_text)
    if not snippet:
        raise RuntimeError("window.__NUXT__ block not found")

    ctx = quickjs.Context()
    ctx.eval("var window = {};")
    ctx.eval(snippet)
    nuxt_json = ctx.eval("JSON.stringify(window.__NUXT__)")
    nuxt = json.loads(nuxt_json)
    detail = None
    breadcrumbs = None
    source = "lcsc"
    data = nuxt.get("data") if isinstance(nuxt, dict) else None
    if isinstance(data, list) and data:
        first = data[0]
        # JLCPCB part detail pages use componentInfo
        if isinstance(first, dict) and "componentInfo" in first:
            source = "jlcpcb"
            detail = first.get("componentInfo")
            breadcrumbs = first.get("navList")
        else:
            detail = first.get("detail") if isinstance(first, dict) else None
            breadcrumbs = first.get("breadcrumbs") if isinstance(first, dict) else None
    return {"detail": detail, "breadcrumbs": breadcrumbs, "source": source}


def _scrape_html_to_json(html_path, json_path):
    html_text = html_path.read_text(encoding="utf-8", errors="ignore")
    if _extract_nuxt_blob(html_text) is not None:
        payload = _parse_nuxt_detail(html_text)
    else:
        payload = _parse_next_detail(html_text)
        if payload is None:
            payload = {"detail": None, "breadcrumbs": None, "note": "product_data_not_found"}
    json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def _extract_model_name_from_json(json_path):
    try:
        data = json.loads(json_path.read_text(encoding="utf-8", errors="ignore"))
    except Exception:
        return None, None
    detail = data.get("detail") if isinstance(data, dict) else None
    if not isinstance(detail, dict):
        return None, None
    # LCSC
    model = detail.get("productModel")
    name = detail.get("productNameEn")
    if model or name:
        return model, name
    # JLCPCB
    model = detail.get("componentModelEn") or detail.get("componentModel")
    name = detail.get("describe") or detail.get("componentNameEn") or detail.get("matchedPartDetail")
    return model, name


def _extract_rohs_from_json(json_path):
    try:
        data = json.loads(json_path.read_text(encoding="utf-8", errors="ignore"))
    except Exception:
        return None
    if not isinstance(data, dict):
        return None
    detail = data.get("detail")
    if not isinstance(detail, dict):
        return None
    source = data.get("source")
    # LCSC: isEnvironment true/false
    if source == "lcsc" or "isEnvironment" in detail:
        val = detail.get("isEnvironment")
        if isinstance(val, bool):
            return "RoHS" if val else "Non-RoHS"
        return None
    # JLCPCB: rohsFlag (1/0)
    if "rohsFlag" in detail:
        val = detail.get("rohsFlag")
        if val == 1 or val is True:
            return "RoHS"
        if val == 0 or val is False:
            return "Non-RoHS"
    return None


def main():
    parser = argparse.ArgumentParser(description="Prepare an AI comparison of project parts and LCSC product data.")
    parser.add_argument("input", nargs="?", help="Legacy standalone .xls, .xlsx, or .csv BOM input")
    parser.add_argument("--project-json", type=Path, help="Project JSON to check, or notes for a legacy BOM input")
    parser.add_argument("--output", type=Path, help="AI check report filename")
    parser.add_argument("--project-directory", type=Path, help="Original project folder for legacy product-cache migration")
    parser.add_argument("--sleep", type=float, default=0.2, help="Seconds to sleep between downloads")
    parser.add_argument(
        "--scrape-only",
        action="store_true",
        help="Do not download; only scrape existing .htm files into .json",
    )
    parser.add_argument(
        "--force-download",
        action="store_true",
        help="Redownload HTML even if it already exists",
    )
    parser.add_argument(
        "--force-scrape",
        action="store_true",
        help="Rescrape HTML into JSON even if JSON already exists",
    )
    args = parser.parse_args()

    if not args.input and not args.project_json:
        parser.error("Provide --project-json or a standalone BOM input.")
    internal = not args.input
    input_path = args.project_json if internal else Path(args.input)
    if not input_path.exists():
        print(f"ERROR: Input file not found: {input_path}", file=sys.stderr)
        return 1

    note_records = {}
    if internal:
        print("CHECK_STAGE: Reading current project…", flush=True)
        try:
            note_records = load_check_notes(input_path)
            rows = project_check_rows(note_records)
        except (OSError, ValueError, KeyError) as exc:
            print(f"ERROR: Invalid project or normalization rules: {exc}", file=sys.stderr)
            return 1
        if not rows:
            print("ERROR: No non-excluded parts to check.")
            return 1
    else:
        ext = input_path.suffix.lower()
        print("CHECK_STAGE: Reading workbook…", flush=True)
        if ext == ".csv":
            rows = _read_csv_rows(input_path, with_designators=True)
        elif ext in (".xls", ".xlsx"):
            rows = _read_excel_rows(input_path, with_designators=True)
        else:
            print(f"ERROR: Unsupported file extension: {ext}", file=sys.stderr)
            return 1

    parts = _normalize_parts([r[1] for r in rows])
    print(f"CHECK_TOTAL: {len(parts)}", flush=True)
    if not parts:
        if not internal:
            print("ERROR: No C-part numbers found in last column.")
            return 1
        print("No assigned LCSC numbers; the report will list unassigned parts.")

    print("CHECK_STAGE: Loading check settings and notes…", flush=True)
    try:
        instruction_lines = load_ai_instructions()
        if not internal and args.project_json:
            note_records = load_check_notes(args.project_json)
    except (OSError, ValueError, KeyError) as exc:
        print(f"ERROR: Invalid AI check settings: {exc}", file=sys.stderr)
        return 1

    out_dir = cache_directory()
    out_dir.mkdir(parents=True, exist_ok=True)
    print(f"Shared LCSC cache: {out_dir}")
    print("CHECK_STAGE: Preparing product cache…", flush=True)
    local_cache = (args.project_directory or input_path.parent) / "lcsc"
    imported, _ = import_local_cache(local_cache)
    if imported:
        print(f"Imported {imported} cached files from {local_cache}")

    downloaded = 0
    skipped = 0
    if not args.scrape_only:
        for part in parts:
            out_path = out_dir / f"{part}.htm"
            try:
                if out_path.exists() and not args.force_download:
                    print(f"{part}: already have HTML")
                    skipped += 1
                    continue
                if _download_html(part, out_path, sleep_s=args.sleep, overwrite=args.force_download):
                    downloaded += 1
            except Exception as exc:
                print(f"Failed {part}: {exc}", file=sys.stderr)

    scraped = 0
    for part in parts:
        html_path = out_dir / f"{part}.htm"
        json_path = out_dir / f"{part}.json"
        if not html_path.exists():
            print(f"{part}: missing HTML, skipping scrape")
            continue
        if json_path.exists() and not (args.force_scrape or args.force_download):
            print(f"{part}: already have JSON")
            continue
        try:
            print(f"{part}: scraping HTML -> JSON")
            _scrape_html_to_json(html_path, json_path)
            scraped += 1
        except Exception as exc:
            print(f"{part}: scrape failed: {exc}", file=sys.stderr)

    # Output verification text file
    print("CHECK_STAGE: Writing check reports…", flush=True)
    verify_path = args.output or (input_path.with_name(input_path.stem + "_AI_Check.txt")
                                 if internal else input_path.with_suffix(".txt"))
    verify_path.parent.mkdir(parents=True, exist_ok=True)
    rohs_path = verify_path.with_suffix(".rohs.txt")
    if input_path.resolve() in {verify_path.resolve(), rohs_path.resolve()}:
        print("ERROR: Report filenames must differ from the input file.", file=sys.stderr)
        return 1
    with open(verify_path, "w", encoding="utf-8", newline="\n") as f:
        f.write("\n".join(instruction_lines) + "\n")
        f.write("FORMAT: <original_description>\\t<productModel>: <productNameEn> // optional notes\n")
        if internal:
            f.write("NOTES: Project descriptions include designators, the selected part's normalized name/value and footprint. Supplier details come from the assigned LCSC number. Project notes may qualify differences. Excluded parts are omitted; unassigned parts are marked for review.\n")
        else:
            f.write("NOTES: Lines are constructed from column 1 (original description) and the JSON fields productModel/productNameEn for the matching LCSC part number in the last column.\n")
        f.write("\n")
        for desc, part, refs in rows:
            part_norm = str(part).strip().upper() if part is not None else ""
            if not part_norm.startswith("C"):
                if internal:
                    f.write(verification_line(desc, "", "(no LCSC part number assigned)",
                                              notes_for_row(note_records, refs, "")))
                continue
            json_path = out_dir / f"{part_norm}.json"
            notes = notes_for_row(note_records, refs, part_norm)
            if json_path.exists():
                model, name = _extract_model_name_from_json(json_path)
                model = model or ""
                name = name or ""
                f.write(verification_line(desc, model, name, notes))
            else:
                f.write(verification_line(desc, part_norm, "(missing json)", notes))

    # Output RoHS text file
    rohs_refs = {"compliant": set(), "non_compliant": set(), "unknown": set()}
    from Shared.references import REFERENCE_RE
    with open(rohs_path, "w", encoding="utf-8", newline="\n") as f:
        f.write("FORMAT: <original_description>\\t<rohs_status>\n")
        f.write("NOTES: RoHS is derived from LCSC isEnvironment or JLCPCB rohsFlag when present.\n")
        f.write("\n")
        for desc, part, refs in rows:
            part_norm = str(part).strip().upper() if part is not None else ""
            designators = set(REFERENCE_RE.findall(refs))
            if not part_norm.startswith("C"):
                if internal:
                    rohs_refs["unknown"].update(designators)
                    f.write(f"{desc}\tNot checked (no LCSC part number)\n")
                continue
            json_path = out_dir / f"{part_norm}.json"
            if json_path.exists():
                rohs = _extract_rohs_from_json(json_path) or "Unknown"
                if rohs == "RoHS":
                    rohs_refs["compliant"].update(designators)
                elif rohs == "Non-RoHS":
                    rohs_refs["non_compliant"].update(designators)
                else:
                    rohs_refs["unknown"].update(designators)
                f.write(f"{desc}\t{rohs}\n")
            else:
                rohs_refs["unknown"].update(designators)
                f.write(f"{desc}\tMissing JSON\n")

    print(f"Done. Downloaded: {downloaded}, Skipped: {skipped}, Scraped: {scraped}, Output: {out_dir}")
    print(f"Wrote verification file: {verify_path}")
    print(f"Wrote RoHS file: {rohs_path}")
    print("CHECK_ROHS: " + json.dumps({key: sorted(refs) for key, refs in rohs_refs.items()}), flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
