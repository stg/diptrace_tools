"""Data and file operations for the DipTrace BOM editor (no GUI dependencies)."""
from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
import tempfile
from copy import deepcopy
from collections import OrderedDict
from datetime import datetime
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape

from lxml import etree

import bom_conversions as conversion
from project_config import project_defaults, validate_filename
from bom_values import SI_PREFIX_EXPONENTS, engineering_magnitude, natural_sort_key, value_sort_key

FIELDS = ("Name", "Value", "Pattern", "LCSC Part #")
JSON_FIELDS = FIELDS + ("Notes",)
ALTERNATE_FIELDS = JSON_FIELDS
REFS = conversion.REFERENCE_RE
PART = re.compile(r"C\d+", re.I)


def names_match(input_name, json_name):
    """Allow JSON suffixes and single-letter lowercase placeholders in input names."""
    if not input_name or not json_name:
        return input_name == json_name
    mask = "".join("[A-Za-z]" if "a" <= char <= "z" else re.escape(char) for char in input_name)
    return re.match(mask, json_name, re.IGNORECASE) is not None


def name_has_added_suffix(input_name, json_name):
    """Compatible extra JSON name characters are a suffix warning."""
    return bool(input_name) and len(json_name) > len(input_name) and names_match(input_name, json_name)


def value_has_added_suffix(input_value, json_value):
    """Added JSON detail warns when every original word is retained, including blanks."""
    original, edited = input_value.casefold().split(), json_value.casefold().split()
    return len(edited) > len(original) and edited[:len(original)] == original


def pattern_difference_is_warning(input_pattern, json_pattern):
    """Known package differences are errors; unknown overrides warn."""
    return not (
        conversion.pattern_has_normalization_rule(input_pattern)
        and conversion.pattern_has_normalization_rule(json_pattern)
    )


def ref_key(ref):
    return conversion.natural_reference_key(ref)


def fingerprint(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def read_schematic(path):
    root = etree.parse(str(path), etree.XMLParser(resolve_entities=False, no_network=True)).getroot()
    if root.tag != "Source" or root.get("Type") != "DipTrace-Schematic":
        raise ValueError("Select a DipTrace schematic XML (.dchxml) file.")
    patterns = {p.get("PatternStyle"): p.findtext("Name", "")
                for p in root.findall("./Library/Library/Patterns/Pattern")}
    components = {c.get("ComponentStyle"): c for c in root.findall("./Library/Components/Component")}
    records = {}
    identities = {}
    for instance in root.findall("./Schematic/Components/Part"):
        component = components.get(instance.get("ComponentStyle"))
        if component is None:
            raise ValueError(f"Unknown component style: {instance.get('ComponentStyle')}")
        definitions = component.findall("Part")
        index = int(instance.get("ComponentPart", "0"))
        definition = definitions[index] if index < len(definitions) else definitions[0]
        if definition.get("PartType", "Normal") != "Normal":
            continue
        ref = instance.findtext("RefDes", "").strip().upper()
        if not ref:
            continue
        pattern = definition.find("Pattern")
        if pattern is None:
            pattern = component.find(".//Pattern")
        style = pattern.get("Style", "") if pattern is not None else ""
        if style and style not in patterns:
            raise ValueError(f"Unknown pattern style {style} for {ref}")
        record = {"Name": instance.findtext("Name", definition.findtext("Name", "")),
                  "Value": instance.findtext("Value", definition.findtext("Value", "")),
                  "Pattern": patterns.get(style, ""), "LCSC Part #": ""}
        identity = (instance.get("ComponentStyle"), instance.get("UpdateId"),
                    instance.get("AllowParts") == "Y", index)
        if ref in records:
            previous = identities[ref]
            multipart = previous[2] and identity[2] and previous[:2] == identity[:2]
            if multipart:
                # DipTrace stores each unit, including the power unit, as an
                # instance. Their displayed Name/Value can differ. The lowest
                # ComponentPart (normally 0) represents the physical BOM part.
                if index >= previous[3]:
                    continue
            elif records[ref] != record:
                raise ValueError(f"Inconsistent schematic parts for {ref}")
        records[ref] = record
        identities[ref] = identity
    if not records:
        raise ValueError("No BOM components found in this schematic.")
    return dict(sorted(records.items(), key=lambda item: ref_key(item[0])))


def record_group_key(record, include_lcsc=True):
    fields = JSON_FIELDS if include_lcsc else FIELDS[:3]
    key = tuple(record.get(field, "") for field in fields)
    if include_lcsc:
        key += (json.dumps(record.get("Alternates", []), ensure_ascii=False, sort_keys=True),
                record.get("Output selection", 0))
    return key


def grouped_rows(records, include_lcsc=True):
    groups = OrderedDict()
    fields = JSON_FIELDS if include_lcsc else FIELDS[:3]
    for ref in sorted(records, key=ref_key):
        row = records[ref]
        key = record_group_key(row, include_lcsc)
        groups.setdefault(key, []).append(ref)
    rows = []
    for key, refs in groups.items():
        row = {"RefDes": ", ".join(refs), **dict(zip(fields, key))}
        record = records[refs[0]]
        if include_lcsc and record.get("Alternates"):
            row["Alternates"] = deepcopy(record["Alternates"])
            row["Output selection"] = record.get("Output selection", 0)
        rows.append(row)
    return rows


def expand_rows(rows):
    records = {}
    for row in rows:
        refs = REFS.findall(str(row.get("RefDes", "")))
        if not refs:
            raise ValueError(f"JSON row has no designators: {row!r}")
        for ref in refs:
            ref = ref.upper()
            if ref in records:
                raise ValueError(f"Duplicate JSON designator: {ref}")
            record = {field: str(row.get(field, "") or "") for field in JSON_FIELDS}
            validate_lcsc(record["LCSC Part #"])
            alternates = row.get("Alternates", [])
            if not isinstance(alternates, list) or any(not isinstance(alt, dict) for alt in alternates):
                raise ValueError(f"Alternates for {ref} must be a list of component objects.")
            record["Alternates"] = [
                {field: str(alt.get(field, "") or "") for field in ALTERNATE_FIELDS}
                for alt in alternates
            ]
            for alternate in record["Alternates"]:
                validate_lcsc(alternate["LCSC Part #"])
            selection = row.get("Output selection", 0)
            if type(selection) is not int or not 0 <= selection <= len(alternates):
                raise ValueError(f"Invalid output component selection for {ref}.")
            record["Output selection"] = selection
            records[ref] = record
    return records


def validate_lcsc(value):
    if value and not PART.fullmatch(value):
        raise ValueError("LCSC number must be blank or C followed by digits (for example C487422).")


def atomic_write(path, content):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, temp = tempfile.mkstemp(prefix=f".{path.name}.", dir=path.parent)
    try:
        with os.fdopen(fd, "wb") as stream:
            stream.write(content)
        os.replace(temp, path)
    finally:
        if os.path.exists(temp):
            os.unlink(temp)


def selected_component(record):
    """Return the primary or alternate chosen for assembly output."""
    selection = record.get("Output selection", 0)
    return record["Alternates"][selection - 1] if selection else record


def output_record(ref, record):
    record = selected_component(record)
    normalized = conversion.normalized_row({"RefDes": ref, **record})
    if normalized["_excluded"]:
        return None
    return {"Comment": conversion._comment(normalized),
            "Footprint": normalized["Pattern"],
            "JLCPCB Part #（optional）": record.get("LCSC Part #", "")}


def workbook_grid(path):
    path = Path(path)
    if path.read_bytes()[:2] == b"PK":
        from openpyxl import load_workbook
        book = load_workbook(BytesIO(path.read_bytes()), data_only=True, read_only=True)
        try:
            sheet = book.active
            sheet.reset_dimensions()
            return list(sheet.iter_rows(values_only=True))
        finally:
            book.close()
    import xlrd
    book = xlrd.open_workbook(str(path))
    sheet = book.sheet_by_index(0)
    return [sheet.row_values(i) for i in range(sheet.nrows)]


def read_assignments(path):
    """Accept simple JLC BOM tables and BOM Matching/Assembly Order exports."""
    grid = workbook_grid(path)
    for row_index, row in enumerate(grid[:60]):
        headers = [str(v or "").strip().casefold() for v in row]
        ref_col = next((i for i, h in enumerate(headers) if h in {"designator", "refdes", "designators"}), None)
        cols = [i for i, h in enumerate(headers)
                if h.startswith("jlcpcb part #") or h in {"lcsc part #", "lcsc part number", "lcsc part", "lcsc"}]
        if ref_col is None or not cols:
            continue
        assignments = {}
        part_col = cols[-1]
        for values in grid[row_index + 1:]:
            raw = str(values[part_col] or "").strip().upper() if part_col < len(values) else ""
            if not PART.fullmatch(raw):
                continue
            refs = REFS.findall(str(values[ref_col] or "")) if ref_col < len(values) else []
            for ref in refs:
                ref = ref.upper()
                if ref in assignments and assignments[ref] != raw:
                    raise ValueError(f"Conflicting LCSC assignments for {ref} in workbook.")
                assignments[ref] = raw
        if not assignments:
            raise ValueError("The workbook has no nonempty LCSC assignments.")
        return assignments
    raise ValueError("No Designator and JLCPCB/LCSC Part # table found in workbook.")


def read_generated(path):
    path = Path(path)
    if not path.exists():
        return {}
    grid = workbook_grid(path)
    if not grid:
        return {}
    headers = [str(v or "").strip() for v in grid[0]]
    if "Designator" not in headers:
        raise ValueError(f"No Designator column in {path}")
    result = {}
    for values in grid[1:]:
        row = dict(zip(headers, (str(v or "") for v in values)))
        for ref in REFS.findall(row.get("Designator", "")):
            result[ref.upper()] = row
    return result


class Project:
    def __init__(self, path):
        self.path = Path(path).resolve()
        self.json_path = self.path.with_suffix(".json")
        self.source_hash = fingerprint(self.path)
        self.source = read_schematic(self.path)
        if self.json_path.exists():
            self.document = json.loads(self.json_path.read_text(encoding="utf-8-sig"))
            if self.document.get("schema_version") != 1:
                raise ValueError("Unsupported sidecar JSON schema; expected schema_version 1.")
            self.records = expand_rows(self.document["rows"])
        else:
            self.document = {**project_defaults(self.path.stem), "schema_version": 1,
                             "source": self.path.name, "rows": []}
            self.records = {ref: {**row, "Notes": "", "Alternates": [], "Output selection": 0}
                            for ref, row in self.source.items()}
            self.save()
        self.saved_state = self.state()
        # New components come from the schematic. Existing JSON values,
        # including edited patterns, remain authoritative on reload.
        for ref, row in self.source.items():
            if ref not in self.records:
                self.records[ref] = {**row, "Notes": "", "Alternates": [], "Output selection": 0}
        self.refresh_outputs()
        self.cpl_report = None

    @property
    def output_dir(self):
        path = Path(self.document.get("output_folder", "./order"))
        return path if path.is_absolute() else self.path.parent / path

    @property
    def bom_xls(self):
        filename = self.document.get("bom_file", f"{self.path.stem}_BOM.xls")
        return self.output_dir / validate_filename(filename, "bom_file", xls=True)

    @property
    def cpl_xls(self):
        filename = self.document.get("cpl_output_file", f"{self.path.stem}_CPL.xls")
        return self.output_dir / validate_filename(filename, "cpl_output_file", xls=True)

    @property
    def ai_check_path(self):
        filename = self.document.get("ai_check_file", f"{self.path.stem}_AI_Check.txt")
        validate_filename(filename, "ai_check_file")
        if Path(filename).suffix.lower() != ".txt":
            raise ValueError("ai_check_file must end in .txt.")
        return self.output_dir / filename

    @property
    def cpl_path(self):
        filename = self.document.get("cpl_file", f"{self.path.stem}_CPL.csv")
        if not isinstance(filename, str) or not filename.strip():
            raise ValueError("cpl_file in the project JSON must be a nonempty filename.")
        path = Path(filename)
        if path.name != filename:
            raise ValueError("cpl_file must be a filename beside the schematic, without a folder path.")
        return self.path.parent / path

    def state(self):
        return json.dumps(self.records, ensure_ascii=False, sort_keys=True)

    @property
    def dirty(self):
        return self.state() != self.saved_state

    def save(self):
        self.document["rows"] = grouped_rows(self.records)
        atomic_write(self.json_path, (json.dumps(self.document, ensure_ascii=False, indent=2) + "\n").encode("utf-8"))
        self.saved_state = self.state()

    def refresh_outputs(self):
        self.outputs = []
        self.output_errors = []
        for path in (self.bom_xls,):
            try:
                self.outputs.append((path, read_generated(path)))
            except Exception as exc:
                self.outputs.append((path, {}))
                self.output_errors.append(f"{path.name}: {exc}")

    def generate(self):
        # Validate placement data and its explicit units before replacing outputs.
        cpl = self.cpl_path
        conversion.validate_cpl_input(cpl)
        conversion.load_cpl_adjustments(conversion.CPL_ADJUSTMENTS_PATH)
        conversion.load_normalization_rules(force=True)
        bom_output, cpl_output = self.bom_xls, self.cpl_xls
        decimal_places = self.document.get("cpl_decimals", 3)
        if type(decimal_places) is not int or not 0 <= decimal_places <= 9:
            raise ValueError("cpl_decimals must be an integer from 0 to 9.")
        if str(bom_output).casefold() == str(cpl_output).casefold():
            raise ValueError("BOM and CPL output filenames must differ.")
        self.save()
        self.output_dir.mkdir(parents=True, exist_ok=True)
        selected = {ref: {field: selected_component(record)[field] for field in FIELDS}
                    for ref, record in self.records.items()}
        bom_rows = conversion.assembly_bom_rows(grouped_rows(selected))
        bom_count = conversion.write_bom_xls(bom_rows, bom_output)
        self.cpl_report = conversion.convert_cpl_for_bom_rows(
            cpl, bom_rows, cpl_output, decimal_places=decimal_places)
        cpl_count = self.cpl_report.rows
        self.refresh_outputs()
        return bom_count, cpl_count

    def sync_schematic(self, errors_only=False, refs=None):
        selected_refs = None if refs is None else set(refs)
        if fingerprint(self.path) != self.source_hash:
            raise ValueError("The schematic changed on disk. Reload before syncing.")
        raw = self.path.read_text(encoding="utf-8")
        start = raw.index("<Schematic")
        prefix, body = raw[:start], raw[start:]
        changes = 0

        def patch_part(match):
            nonlocal changes
            text = match.group(0)
            node = etree.fromstring(text.encode("utf-8"))
            ref = node.findtext("RefDes", "").strip().upper()
            if selected_refs is not None and ref not in selected_refs:
                return text
            if ref not in self.source or ref not in self.records:
                return text
            for field in ("Name", "Value"):
                desired = self.records[ref][field]
                current = self.source[ref][field]
                # All retains the existing behavior: only fields changed from
                # the main symbol are written, preserving unchanged unit text.
                if desired == current:
                    continue
                if errors_only and (
                    (field == "Value" and value_has_added_suffix(current, desired))
                    or (field == "Name" and name_has_added_suffix(current, desired))
                ):
                    continue
                if node.findtext(field, "") == desired:
                    continue
                replacement = f"<{field}>{escape(desired)}</{field}>"
                # Direct Name/Value fields precede LibPath, markings, and pins.
                boundary = re.search(r"<(?:LibPath|RefDesMarking|Pins|Category)\b", text)
                end = boundary.start() if boundary else text.index("</Part>")
                head, tail = text[:end], text[end:]
                expression = rf"<{field}(?:\s*/>|>.*?</{field}>)"
                head, n = re.subn(expression, lambda _: replacement, head, count=1, flags=re.S)
                if not n:
                    head += replacement + "\n"
                text = head + tail
                changes += 1
            return text

        body = re.sub(r'<Part\b[^>]*\bId="[^"]*"[^>]*>.*?</Part>', patch_part, body, flags=re.S)
        updated = (prefix + body).encode("utf-8")
        # Validate the complete XML before replacing the schematic.
        etree.fromstring(updated, etree.XMLParser(resolve_entities=False, no_network=True))
        backup = None
        if changes:
            backup = self.path.with_name(f"{self.path.name}.{datetime.now():%Y%m%d-%H%M%S-%f}.bak")
            shutil.copy2(self.path, backup)
            # Preserve Windows line endings from the original file.
            if b"\r\n" in self.path.read_bytes():
                updated = updated.replace(b"\n", b"\r\n")
            atomic_write(self.path, updated)
        self.source = read_schematic(self.path)
        self.source_hash = fingerprint(self.path)
        self.save()
        return changes, backup
